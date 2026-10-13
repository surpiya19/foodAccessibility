# FoodPrint Score - step 2: extract Richmond from the USDA Food Access Research Atlas (FARA)
#
# FARA has no API, so download the 2025 release by hand:
#   https://www.ers.usda.gov/data-products/food-access-research-atlas
#   -> "Download the data" -> save as data/raw/2025-sram-fara-data.xlsx
# The 2025 release uses 2020 census tracts, so all 75 Richmond tracts match
# PLACES and ACS (the 2019 release used 2010 tracts and only 57 matched).
#
# Reads:   data/raw/2025-sram-fara-data.xlsx  (~110 MB, takes a couple of minutes)
# Writes:  data/processed/fara2025_richmond.csv  (75 rows, all three data sheets joined)

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
PROCESSED = ROOT / "data" / "processed"

COUNTY_FIPS_FULL = "51760"  # Richmond city, VA
FARA_XLSX = RAW / "2025-sram-fara-data.xlsx"
SHEETS = ["GeneralTractCharacteristic Data",  # population, poverty, SNAP, vehicles
          "Straight-Line Distance Data",      # SD_SRAM_* low-access measures
          "Driving Distance Data"]            # DD_SRAM_* low-access measures (used in the score)

if not FARA_XLSX.exists():
    raise SystemExit(f"Missing {FARA_XLSX}. Download it from USDA ERS (see top of this file).")

def read_richmond_rows(workbook, sheet):
    # read row by row and keep Richmond only (pd.read_excel loads all ~84k tracts and is slow)
    rows = workbook[sheet].iter_rows(values_only=True)
    next(rows)            # row 0 is a title
    header = next(rows)   # row 1 holds the column names
    keep = [r for r in rows if r[0] and str(r[0]).startswith(COUNTY_FIPS_FULL)]
    df = pd.DataFrame(keep, columns=header)
    df["CensusTract20"] = df["CensusTract20"].astype(str)
    return df


workbook = load_workbook(FARA_XLSX, read_only=True)
parts = []
for sheet in SHEETS:
    print("Reading", sheet, "...")
    df = read_richmond_rows(workbook, sheet)
    if parts:
        df = df.drop(columns=["State", "County20", "County24"])  # already in the first sheet
    parts.append(df)

fara = parts[0]
for df in parts[1:]:
    fara = fara.merge(df, on="CensusTract20")

print("Richmond tracts:", len(fara), "(expect 75)")
PROCESSED.mkdir(parents=True, exist_ok=True)
fara.to_csv(PROCESSED / "fara2025_richmond.csv", index=False)
